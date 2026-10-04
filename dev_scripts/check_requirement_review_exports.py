"""Independent read-only pagination export checks; private QA artifacts stay ignored."""
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
root = Path(__file__).resolve().parents[1] / 'output' / 'playwright' / 'requirement-review'
statuses = {'pending':'待复核','unknown':'未知/待澄清','conflict':'候选冲突','stale':'来源已失效','confirmed':'人工已确认'}
responses = {'pending':'待应答','met':'人工复核：满足','deviation':'人工复核：偏离','not-applicable':'人工复核：不适用'}
for variant in ['dev', 'build-subpath']:
    directory = root / variant
    report = json.loads((directory / 'browser-result.json').read_text(encoding='utf-8'))
    assert report['passed'] and report['maxMs'] <= 5000 and not report['errors'] and not report['external'] and not report['writes']
    pack = json.loads((directory / 'real-working-after-qa.json').read_text(encoding='utf-8'))['package']
    assert len(pack['requirements']) == 1196 and len(pack['sources']) == 5
    sources = {source['id']:source for source in pack['sources']}
    equipment = {item['id']:item['label'] for item in pack['equipment']}
    expected = []
    for requirement in pack['requirements']:
        for candidate in requirement['candidates']:
            evidence = candidate['evidence']; source = sources[evidence['sourceId']]
            assert source['text'][evidence['start']:evidence['end']] == evidence['quote']
            assert any(evidence['start'] >= s['start'] and evidence['end'] <= s['end'] for s in source['segments'])
            expected.append((equipment.get(requirement['equipmentId'], '待分配'), requirement['field'], statuses[requirement['reviewStatus']], candidate['value'] if candidate['value'] is not None else '未知', requirement['confirmedValue'] if requirement['confirmedValue'] is not None else '未知', '未核对' if requirement['forbiddenDeviation'] is None else '是' if requirement['forbiddenDeviation'] else '否', responses[requirement['response']], evidence['quote']))
    book = load_workbook(directory / 'real-requirements.xlsx')
    rows = list(book['技术要求工作稿'].values)
    assert len(rows) == len(expected) + 2 and rows[0][1] == f"版本{pack['revision']}"
    assert [tuple(row[:8]) for row in rows[2:]] == expected
    assert book['修订记录'].max_row == len(pack['revisions']) + 1, '修订导出被截断为20条'
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row), '导出含执行公式'
    assert 'QA未记录值不得进入导出' not in json.dumps(rows, ensure_ascii=False)
    synthetic = json.loads((directory / 'synthetic-backup.json').read_text(encoding='utf-8'))['package']
    assert '未记录草稿A' not in json.dumps(synthetic, ensure_ascii=False)
    assert len(synthetic['requirements']) == 46
    print(f'{variant}: 全量1196要求/证据/修订与JSON一致，草稿未混入，静态Excel无公式；最大{report["maxMs"]}ms')
