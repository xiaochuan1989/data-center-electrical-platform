"""只读核验隔离浏览器合成导出；不创建或改写客户工作簿。"""
from pathlib import Path
from zipfile import ZipFile
import json
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

root = Path('output/playwright/nonstandard-n3')
expected = ['方案概览', '配置清单', '采用配置', '技术要求与应答', '澄清与边界', '修订记录']
books = {}
for name in ['cabinet', 'cabinet-revised', 'pdu-draft', 'busway-draft', 'requirements']:
    with ZipFile(root / f'{name}.xlsx') as archive:
        assert archive.testzip() is None, 'CRC32失败'
        assert not any('vba' in path.lower() or 'externalLink' in path for path in archive.namelist())
        for path in archive.namelist():
            if path.endswith('.xml') or path.endswith('.rels'):
                ET.fromstring(archive.read(path))
    book = load_workbook(root / f'{name}.xlsx', data_only=False)
    assert book.sheetnames == (['技术要求工作稿', '修订记录'] if name == 'requirements' else expected)
    assert not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)
    books[name] = book

cabinet = books['cabinet']['配置清单']
assert cabinet['G2'].value == 4 and cabinet['G3'].value == 2
assert cabinet['C3'].value == '=1+1' and cabinet['C3'].data_type == 's'
assert cabinet.freeze_panes == 'A2'
assert books['pdu-draft']['配置清单']['G2'].value == '待确认'
requirements = books['requirements']['技术要求工作稿']
assert requirements['A3'].value == 'TEST-柜'
assert requirements['E3'].value == 630
assert requirements['G3'].value == '人工复核：满足'
assert '630A' in requirements['H3'].value
assert requirements['I3'].value
snapshot = json.loads((root / 'working.json').read_text(encoding='utf-8'))
scheme = snapshot['package']['schemes'][0]
equipment = snapshot['package']['equipment'][0]
revised = books['cabinet-revised']['配置清单']
for index, component in enumerate(scheme['components'], 2):
    assert revised.cell(index, 5).value == component['quantity']
    assert revised.cell(index, 6).value == equipment['quantity']
    assert revised.cell(index, 7).value == component['quantity'] * equipment['quantity']
    assert revised.cell(index, 17).value == snapshot['packageRevision']
assert '<script>' not in (root / 'cabinet.html').read_text(encoding='utf-8')
print('PASS: 5 XLSX CRC/XML, typed quantities, unknowns, evidence, revision, no formulas/macros/links, HTML escaping')
