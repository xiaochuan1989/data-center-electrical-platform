#!/usr/bin/env python3
"""从当前有效工程工具中提取非价格型基础数据，供静态网页只读使用。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / "工具模板"
DATA = ROOT / "src" / "data"


def find_one(name: str) -> Path:
    matches = list(TOOLS.rglob(name))
    if len(matches) != 1:
        raise RuntimeError(f"预期唯一文件 {name}，实际找到 {len(matches)} 个")
    return matches[0]


def clean(value):
    if isinstance(value, str):
        return value.strip()
    return value


def write_json(name: str, records: list[dict]) -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    output = DATA / name
    output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{name}: {len(records)} 条")


def extract_busbars() -> list[dict]:
    workbook = load_workbook(find_one("铜排载流量-A03.xlsx"), data_only=True, read_only=True)
    sheet = workbook["Data"]
    result = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        spec, configuration, coated, bare, area, note = [clean(row[i]) if len(row) > i else None for i in range(1, 7)]
        if not spec or not isinstance(coated, (int, float)):
            continue
        result.append({
            "spec": str(spec),
            "configuration": configuration or "",
            "coatedCurrentA": coated,
            "bareCurrentA": bare if isinstance(bare, (int, float)) else None,
            "areaMm2": area if isinstance(area, (int, float)) else None,
            "note": note or ""
        })
    workbook.close()
    return result


def extract_conductors_and_cables() -> tuple[list[dict], list[dict], list[dict]]:
    workbook = load_workbook(find_one("A-导体电缆选型-A05.xlsx"), data_only=True, read_only=True)
    conductors = []
    for row in workbook["导体数据库"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not isinstance(row[4] if len(row) > 4 else None, (int, float)):
            continue
        conductors.append({
            "type": clean(row[0]),
            "parallelCount": clean(row[1]),
            "orientation": clean(row[2]),
            "ambientC": clean(row[3]),
            "current70A": clean(row[4]),
            "current105A": clean(row[5]),
            "size": clean(row[7]) if len(row) > 7 else ""
        })

    cables = []
    for row in workbook["电缆数据库"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not isinstance(row[4] if len(row) > 4 else None, (int, float)):
            continue
        cables.append({
            "type": clean(row[0]),
            "coreCount": clean(row[1]),
            "installation": clean(row[2]),
            "ambientC": clean(row[3]),
            "currentA": clean(row[4]),
            "size": clean(row[6]) if len(row) > 6 else "",
            "arrangement": clean(row[7]) if len(row) > 7 else ""
        })

    awg = []
    awg_sheet = next((sheet for sheet in workbook.worksheets if "线规" in sheet.title or "AWG" in sheet.title.upper()), None)
    if awg_sheet:
        for row in awg_sheet.iter_rows(values_only=True):
            cells = [clean(value) for value in row]
            if not cells or not isinstance(cells[0], (int, float, str)):
                continue
            text = " ".join(str(value) for value in cells if value is not None)
            if text and any(isinstance(value, (int, float)) for value in cells):
                awg.append({"values": cells})
    workbook.close()
    return conductors, cables, awg


def main() -> int:
    busbars = extract_busbars()
    conductors, cables, awg = extract_conductors_and_cables()
    write_json("busbar-catalog.json", busbars)
    write_json("conductor-catalog.json", conductors)
    write_json("cable-catalog.json", cables)
    write_json("awg-catalog.json", awg)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
