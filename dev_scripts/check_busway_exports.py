"""Independent read-only checks of synthetic regional busway browser exports."""
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

root = Path('output/playwright/busway-assistant')
working = json.loads((root / 'busway-working.json').read_text(encoding='utf-8'))
invalid = json.loads((root / 'busway-invalid-working.json').read_text(encoding='utf-8'))
scheme = working['package']['schemes'][0]
assert scheme['origin'] == 'busway-assistant'
assert scheme['assistant']['quantityBasis'] == 'region'
assert working['package']['equipment'][0]['quantity'] == 1
assert invalid['package']['equipment'][0]['quantity'] == 2
assert scheme['config']['endpointCount'] == 13 and scheme['config']['buswayLengthM'] == 130.5
assert scheme['config']['terminalCount'] is None and scheme['config']['plugBoxOutputs'] is None
assert all(r['reviewStatus'] != 'confirmed' for r in working['package']['requirements'])
for stem in ['busway-draft', 'busway-revised', 'busway-invalid-basis']:
    with ZipFile(root / f'{stem}.xlsx') as archive:
        assert archive.testzip() is None
        assert not any('vba' in p.lower() or 'externalLink' in p for p in archive.namelist())
        for p in archive.namelist():
            if p.endswith(('.xml', '.rels')):
                ET.fromstring(archive.read(p))
    book = load_workbook(root / f'{stem}.xlsx', data_only=False)
    assert len(book.sheetnames) == 6
    assert not any(c.data_type == 'f' for sheet in book for row in sheet for c in row)
    assert '人工区域供货表' in book['方案概览']['B3'].value
    rows = list(book['配置清单'].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 5 and rows[-1][6] == '待确认'
    if stem == 'busway-draft':
        assert [r[6] for r in rows] == [3, 125.5, 12, 64, '待确认']
    else:
        snapshot = invalid if stem == 'busway-invalid-basis' else working
        for row, component in zip(rows, snapshot['package']['schemes'][0]['components']):
            assert row[2] == component['name']
            assert row[4] == (component['quantity'] if component['quantity'] is not None else '待确认')
            assert row[16] == snapshot['packageRevision']
            if stem == 'busway-invalid-basis':
                assert row[6] == '待确认', '区域口径错误时不得翻倍'
            else:
                assert row[6] == (component['quantity'] if component['quantity'] is not None else '待确认')
        notes = '\n'.join(str(c.value or '') for row in book['澄清与边界'] for c in row)
        assert '长度130.5 m' in notes and '端口箱13' in notes
        assert '三相四线制' in notes and '暂定' in notes
html = (root / 'busway-revised.html').read_text(encoding='utf-8')
assert '长度130.5 m' in html and '端口箱13' in html and '<script>' not in html
result = json.loads((root / 'browser-result.json').read_text(encoding='utf-8'))
assert result['passed'] and result['synthetic'] and result['invalidBasisNoMultiply']
assert not result['errors'] and not result['externalWrites']
print('PASS: 3 regional XLSX CRC/XML, six sheets, same revision/quantities, invalid basis null totals, dynamic HTML, no formulas/macros/links')
