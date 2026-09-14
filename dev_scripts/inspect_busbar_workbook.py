import json
import sys
from pathlib import Path

import openpyxl


sys.stdout.reconfigure(encoding="utf-8")
root = Path(__file__).resolve().parents[1]
path = root / "工具模板" / "铜排载流量-A03.xlsx"


def cell_dump(workbook, sheet_name, min_row, max_row, min_col, max_col):
    sheet = workbook[sheet_name]
    rows = []
    for row in sheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        values = []
        for cell in row:
            value = cell.value
            if isinstance(value, openpyxl.worksheet.formula.ArrayFormula):
                value = {"array_formula": value.text, "ref": value.ref}
            values.append(value)
        if any(value is not None for value in values):
            rows.append({"row": row[0].row, "values": values})
    return rows


formula_wb = openpyxl.load_workbook(path, data_only=False)
value_wb = openpyxl.load_workbook(path, data_only=True)

result = {
    "path": str(path),
    "sheets": formula_wb.sheetnames,
    "defined_names": [
        {
            "name": item.name,
            "attr_text": item.attr_text,
            "hidden": item.hidden,
        }
        for item in formula_wb.defined_names.values()
    ],
    "calc_formulas": cell_dump(formula_wb, "Calc", 1, 24, 1, 6),
    "calc_cached_values": cell_dump(value_wb, "Calc", 1, 24, 1, 6),
    "data_validations": [],
}

for validation in formula_wb["Calc"].data_validations.dataValidation:
    result["data_validations"].append(
        {
            "sqref": str(validation.sqref),
            "type": validation.type,
            "formula1": validation.formula1,
            "formula2": validation.formula2,
        }
    )

print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
