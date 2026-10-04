"""Independent, read-only validation of synthetic PDU browser exports."""
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

root = Path('output/playwright/pdu-assistant')
snapshot = json.loads((root / 'pdu-working.json').read_text(encoding='utf-8'))
scheme = snapshot['package']['schemes'][0]
equipment = snapshot['package']['equipment'][0]
for stem in ['pdu-draft', 'pdu-revised']:
    with ZipFile(root / f'{stem}.xlsx') as z:
        assert z.testzip() is None
        assert not any('vba' in p.lower() or 'externalLink' in p for p in z.namelist())
        for p in z.namelist():
            if p.endswith(('.xml', '.rels')):
                ET.fromstring(z.read(p))
    book = load_workbook(root / f'{stem}.xlsx', data_only=False)
    assert len(book.sheetnames) == 6
    assert not any(c.data_type == 'f' for sheet in book for row in sheet for c in row)
    assert book['配置清单']['G2'].value == '待确认'
    assert '原文条款候选' in book['方案概览']['B3'].value
    if stem == 'pdu-revised':
        rows = list(book['配置清单'].iter_rows(min_row=2, values_only=True))
        for row, component in zip(rows, scheme['components']):
            assert row[2] == component['name']
            assert row[4] == (component['quantity'] if component['quantity'] is not None else '待确认')
            total = component['quantity'] * equipment['quantity'] if component['quantity'] is not None else '待确认'
            assert row[6] == total and row[16] == snapshot['packageRevision']
        notes = '\n'.join(str(c.value or '') for row in book['澄清与边界'] for c in row)
        assert '插座15位' in notes and '后续' in notes
assert scheme['config']['outputCount'] == 15
assert scheme['origin'] == 'pdu-assistant'
assert all(r['reviewStatus'] != 'confirmed' for r in snapshot['package']['requirements'])
html = (root / 'pdu-revised.html').read_text(encoding='utf-8')
assert '插座15位' in html and '后续设计联络' in html and '<script>' not in html
result = json.loads((root / 'browser-result.json').read_text(encoding='utf-8'))
assert result['passed'] and result['synthetic'] and not result['errors'] and not result['externalWrites']
print('PASS: 2 XLSX CRC/XML, 6 sheets, unknowns, JSON quantities/revision, dynamic HTML, pending states, no formulas/macros/links')
